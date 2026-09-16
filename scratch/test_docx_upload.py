"""DOCX resume uploads: extraction, and the upload routes accepting them.

The landing page advertised DOCX while every route refused it with "Please
upload a PDF or TXT file". These checks build real .docx files in memory with
python-docx -- including the two-column TABLE layout many Word resume
templates use, where reading paragraphs alone returns almost nothing.

    python scratch/test_docx_upload.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document  # noqa: E402

from backend.services import helpers  # noqa: E402

PASSED, FAILED = [], []


def check(name, cond, detail=""):
    (PASSED if cond else FAILED).append(name)
    print("  %s  %s%s" % ("PASS" if cond else "FAIL", name, "" if cond else "  " + str(detail)))


class Upload(io.BytesIO):
    def __init__(self, data, filename):
        super().__init__(data)
        self.filename = filename


def docx_bytes(build):
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def plain(doc):
    doc.add_heading("Priya Sharma", 0)
    doc.add_paragraph("Software Engineer | priya@example.com")
    doc.add_heading("Experience", 1)
    doc.add_paragraph("Built a Flask API serving 40k requests a day")


def tabled(doc):
    doc.add_paragraph("Arjun Mehta")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "SKILLS\nPython\nSQL"
    t.cell(0, 1).text = "EXPERIENCE\nData Analyst at Acme\nCut report time by half"
    doc.add_paragraph("References available on request")


def main():
    print("\n[1] Extraction")
    text = helpers.extract_resume_text(Upload(docx_bytes(plain), "cv.docx"))
    check("paragraphs are extracted", "Priya Sharma" in text and "Flask API" in text, text)
    text = helpers.extract_resume_text(Upload(docx_bytes(tabled), "Resume.DOCX"))
    check("table cells are extracted", "Python" in text and "Data Analyst at Acme" in text, text)
    check("document order is kept",
          text.index("Arjun Mehta") < text.index("SKILLS") < text.index("References"), text)
    check("a corrupt .docx reads as empty, not a crash",
          helpers.extract_resume_text(Upload(b"not a zip", "bad.docx")) == "")
    check("TXT still decodes", helpers.extract_resume_text(Upload(b"hello", "a.txt")) == "hello")
    check("accepted extensions", all(helpers.is_resume_filename(n) for n in ["a.pdf", "a.DOCX", "a.txt"]))
    check("legacy .doc is still refused", not helpers.is_resume_filename("a.doc"))
    check("docx mime type", helpers.resume_mime_type("x.docx").endswith("wordprocessingml.document"))

    print("\n[2] Routes accept DOCX")
    from backend import create_app
    app = create_app()
    app.config["TESTING"] = True
    client = app.test_client()
    # An unreadable .docx must get past the extension gate and fail on
    # CONTENT -- proving the route no longer refuses the format itself,
    # without spending a real AI call.
    r = client.post("/api/analyze", data={"resume": (io.BytesIO(b"junk"), "cv.docx")},
                    content_type="multipart/form-data")
    body = r.get_json() or {}
    check("/api/analyze no longer rejects the .docx extension",
          r.status_code == 400 and "PDF, DOCX or TXT file" not in body.get("error", ""), (r.status_code, body))
    r = client.post("/api/analyze", data={"resume": (io.BytesIO(b"x"), "cv.doc")},
                    content_type="multipart/form-data")
    check("/api/analyze names DOCX when refusing a format",
          r.status_code == 400 and "DOCX" in (r.get_json() or {}).get("error", ""), r.get_json())

    print("\n%d passed, %d failed" % (len(PASSED), len(FAILED)))
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
